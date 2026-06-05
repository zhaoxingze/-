# -*- coding: utf-8 -*-
from pathlib import Path
import csv

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
REPORT = ROOT / "report" / "final_report.docx"

with (RESULTS / "metrics.csv").open("r", encoding="utf-8-sig", newline="") as f:
    ROWS = list(csv.DictReader(f))


def set_font(run, size=11, bold=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_para(doc, text="", first=True, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    if first:
        p.paragraph_format.first_line_indent = Cm(0.74)
    if text:
        set_font(p.add_run(text))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading("", level=level)
    color = (46, 116, 181) if level <= 2 else (31, 77, 120)
    set_font(p.add_run(text), 16 if level == 1 else 13, True, color)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)


def shade_cell(cell, fill="F2F4F7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=8.5):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    set_font(p.add_run(text), size, bold)


def set_table_widths(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for i, width in enumerate(widths_cm):
            row.cells[i].width = Cm(width)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(text), 10, False, (90, 90, 90))
    p.paragraph_format.space_after = Pt(8)


def add_results_table(doc):
    table = doc.add_table(rows=1, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["编号", "方法", "准确率", "标准差", "损失", "可训练参数", "可训练比例"]
    widths = [1.0, 2.7, 1.4, 1.3, 1.2, 2.2, 1.5]
    for i, header in enumerate(headers):
        set_cell(table.rows[0].cells[i], header, True)
        shade_cell(table.rows[0].cells[i])
    for row in ROWS:
        cells = table.add_row().cells
        set_cell(cells[0], row["experiment_id"])
        set_cell(cells[1], row["method"])
        set_cell(cells[2], f"{float(row['test_acc_mean']) * 100:.2f}%")
        set_cell(cells[3], f"{float(row['test_acc_std']) * 100:.2f}%")
        set_cell(cells[4], f"{float(row['test_loss_mean']):.4f}")
        set_cell(cells[5], f"{int(row['trainable_params']):,}")
        set_cell(cells[6], f"{float(row['trainable_percent']):.4f}%")
    set_table_widths(table, widths)


def add_image(doc, path, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(5.8))
    add_caption(doc, text)


def pct(row):
    return float(row["test_acc_mean"]) * 100


def trainable(row):
    return float(row["trainable_percent"])


def build():
    main = {r["experiment_id"]: r for r in ROWS if r["experiment_id"] in ["E1", "E2", "E3", "E4"]}
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(11)

    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("技术报告"), 22, True)
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("基于 ImageNet 预训练 ResNet18 的高效微调方法探索"), 17, True)
    for _ in range(4):
        doc.add_paragraph()
    for label in ["姓名", "学号", "班级", "指导教师"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(f"{label}：___________________"), 12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("完成日期：2026 年 6 月"), 12)
    doc.add_page_break()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("基于 ImageNet 预训练 ResNet18 的高效微调方法探索"), 18, True)

    add_heading(doc, "摘要")
    add_para(
        doc,
        "本文围绕课程设计“模型微调技术探索”要求，基于 ImageNet 预训练 ResNet18 和 CIFAR-10 小样本分类任务，"
        "比较 Linear Probing、Partial Fine-tuning、Full Fine-tuning、SSF-All，并补充 SSF 层级、小样本比例、LoRA 和 Adapter 实验。"
        "实验使用每类 100 张训练图像与 100 张测试图像，训练 4 个 epoch，并在两个随机种子上取平均。",
    )
    add_para(
        doc,
        f"结果显示，E2 Partial Fine-tuning 在完整小样本设置下准确率最高，为 {pct(main['E2']):.2f}%；"
        f"E4 SSF-All 只训练 {trainable(main['E4']):.4f}% 参数，准确率达到 {pct(main['E4']):.2f}%，体现出更好的参数效率。"
        "LoRA 在本实验中只作用于新初始化分类头，结果应作为简化实现局限分析。",
    )
    add_para(doc, "关键词：模型微调；参数高效微调；ResNet18；SSF；LoRA；Adapter；CIFAR-10", False)

    add_heading(doc, "1 引言")
    add_para(doc, "预训练模型已经成为视觉识别任务的重要基础。相比从零训练，使用 ImageNet 预训练模型可以显著降低数据需求，但全参数微调仍然会带来计算成本、存储成本和小样本过拟合风险。")
    add_para(doc, "参数高效微调通过冻结大部分骨干参数，只训练少量任务相关参数完成迁移。SSF、LoRA 和 Adapter 分别从特征缩放平移、低秩增量和瓶颈残差修正角度提供了不同实现路径。")

    add_heading(doc, "2 方法与实验设计")
    add_para(doc, "主体实验包括 E1 Linear Probing、E2 Partial Fine-tuning、E3 Full Fine-tuning 和 E4 SSF-All。E5/E6 研究 SSF 只作用于高层或较少层时的效果；E7/E8 分析训练数据减少到 25% 和 50% 时不同方法的变化。LoRA 和 Adapter 作为保留扩展实验。")
    add_para(doc, "所有实验均采用 ImageNet 预训练 ResNet18。E1 只训练分类头；E2 解冻 layer4 和分类头；E3 更新全部参数；E4 在多个 ResNet stage 后加入 scale/shift 参数并训练分类头。")

    add_heading(doc, "3 实验结果")
    add_results_table(doc)
    add_caption(doc, "表 1 新实验结果汇总")
    add_image(doc, RESULTS / "accuracy_comparison.jpg", "图 1 不同方法准确率对比")
    add_image(doc, RESULTS / "parameter_comparison.jpg", "图 2 可训练参数比例对比")
    add_image(doc, RESULTS / "training_curves.jpg", "图 3 测试准确率训练曲线")

    add_heading(doc, "4 结果分析")
    add_para(doc, f"E1 Linear Probing 准确率为 {pct(main['E1']):.2f}%，可训练参数比例仅 {trainable(main['E1']):.4f}%，说明预训练特征本身已经具备较强迁移能力。")
    add_para(doc, f"E2 Partial Fine-tuning 准确率最高，但可训练参数比例达到 {trainable(main['E2']):.4f}%，参数效率不如 SSF。E3 Full Fine-tuning 更新全部参数，但准确率为 {pct(main['E3']):.2f}%，说明小样本下全量更新未必更优。")
    add_para(doc, f"E4 SSF-All 只训练 {trainable(main['E4']):.4f}% 参数，准确率为 {pct(main['E4']):.2f}%，在性能与参数量之间取得较好折中。E5/E6 表明减少 SSF 作用层后仍能保持接近效果，说明高层特征调整对迁移分类尤为重要。")
    add_para(doc, "E7 和 E8 显示训练数据减少后，所有方法准确率均下降；局部解冻仍然保持较强表现，但参数量较高。LoRA 的低准确率主要源于它只作用于新分类头，没有作用于预训练骨干层。Adapter 的结果接近 Linear Probing，说明仅在最终特征后添加瓶颈模块的表达空间仍有限。")

    add_heading(doc, "5 结论")
    add_para(doc, "本文完成了课程设计要求的高效微调实验与分析。综合结果看，Partial Fine-tuning 在当前设置下取得最高准确率，SSF-All 在极低可训练参数比例下取得较好性能，更符合参数高效微调主题。后续可进一步把 LoRA 与 Adapter 插入预训练骨干内部，并增加训练轮数与随机种子以提高结论稳定性。")

    add_heading(doc, "参考文献")
    for i, ref in enumerate(
        [
            "Chen S. et al. AdaptFormer: Adapting Vision Transformers for Scalable Visual Recognition. NeurIPS, 2022.",
            "Lian D. et al. Scaling & Shifting Your Features: A New Baseline for Efficient Model Tuning. NeurIPS, 2022.",
            "Hu E. J. et al. LoRA: Low-Rank Adaptation of Large Language Models. ICLR, 2022.",
            "Ben Zaken E. et al. BitFit: Simple Parameter-efficient Fine-tuning for Transformer-based Masked Language-models. ACL, 2022.",
            "He K. et al. Deep Residual Learning for Image Recognition. CVPR, 2016.",
        ],
        1,
    ):
        add_para(doc, f"[{i}] {ref}", False)

    doc.save(REPORT)
    print(REPORT, REPORT.stat().st_size)


if __name__ == "__main__":
    build()
