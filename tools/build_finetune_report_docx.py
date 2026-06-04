from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"F:\Exploration")
OUT_DIR = ROOT / "outputs" / "finetune"
FINAL_DOCX = ROOT / "模型微调技术探索_课程设计小论文.docx"
CURVE_PNG = OUT_DIR / "finetune_accuracy_curves.png"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def set_east_asian_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_paragraph_format(paragraph, first_line: bool = False) -> None:
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(6)
    if first_line:
        paragraph.paragraph_format.first_line_indent = Cm(0.74)


def add_para(doc: Document, text: str = "", *, first_line: bool = True, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    set_paragraph_format(p, first_line=first_line)
    if text:
        r = p.add_run(text)
        set_east_asian_font(r, 11)
    return p


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading("", level=level)
    r = p.add_run(text)
    size = 16 if level == 1 else 13 if level == 2 else 12
    set_east_asian_font(r, size, True)
    r.font.color.rgb = RGBColor(46, 116, 181) if level <= 2 else RGBColor(31, 77, 120)
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)


def add_numbered_items(doc: Document, items: list[str]) -> None:
    for idx, item in enumerate(items, 1):
        p = doc.add_paragraph(style="List Number")
        set_paragraph_format(p, first_line=False)
        r = p.add_run(item)
        set_east_asian_font(r, 11)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text: str, *, bold: bool = False, align=WD_ALIGN_PARAGRAPH.CENTER) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_east_asian_font(r, 9.5, bold)


def set_table_widths(table, widths_cm: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for i, width in enumerate(widths_cm):
            row.cells[i].width = Cm(width)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_east_asian_font(r, 10)
    r.font.color.rgb = RGBColor(85, 85, 85)


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for style_name in ["List Number", "List Bullet"]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
        style.font.size = Pt(11)


def add_cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("技术报告")
    set_east_asian_font(r, 22, True)

    for _ in range(2):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("基于 ImageNet 预训练模型的参数高效微调方法探索")
    set_east_asian_font(r, 17, True)

    for _ in range(4):
        doc.add_paragraph()

    for label in ["姓名", "学号", "班级", "指导教师"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{label}：____________________")
        set_east_asian_font(r, 12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("完成日期：2026 年 6 月")
    set_east_asian_font(r, 12)
    doc.add_page_break()


def add_result_table(doc: Document, summary: list[dict[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["方法", "平均准确率", "标准差", "测试损失", "平均耗时/s", "可训练参数", "可训练比例"]
    widths = [2.0, 2.1, 1.6, 1.7, 1.8, 2.1, 1.9]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True)
        shade_cell(table.rows[0].cells[i], "F2F4F7")
    for row in summary:
        cells = table.add_row().cells
        set_cell_text(cells[0], row["method"])
        set_cell_text(cells[1], f"{float(row['test_acc_mean']) * 100:.2f}%")
        set_cell_text(cells[2], f"{float(row['test_acc_std']) * 100:.2f}%")
        set_cell_text(cells[3], f"{float(row['test_loss_mean']):.4f}")
        set_cell_text(cells[4], f"{float(row['time_sec_mean']):.2f}")
        set_cell_text(cells[5], f"{int(float(row['trainable_params'])):,}")
        set_cell_text(cells[6], f"{float(row['trainable_percent']):.4f}%")
    set_table_widths(table, widths)


def add_method_table(doc: Document) -> None:
    rows = [
        ["LINEAR", "冻结骨干网络，只训练分类头", "分类头参数"],
        ["BITFIT", "训练偏置项和分类头", "偏置项 + 分类头"],
        ["SSF", "在特征上学习缩放和平移", "scale/shift + 分类头"],
        ["LORA", "用低秩矩阵近似分类头增量", "低秩矩阵 + 偏置"],
        ["ADAPTER", "添加瓶颈适配器进行特征变换", "适配器 + 分类头"],
        ["FULL", "更新全部模型参数", "全部参数"],
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["方法", "实验实现", "可训练部分"]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True)
        shade_cell(table.rows[0].cells[i], "F2F4F7")
    for row in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], row[0])
        set_cell_text(cells[1], row[1], align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(cells[2], row[2], align=WD_ALIGN_PARAGRAPH.LEFT)
    set_table_widths(table, [2.1, 6.8, 4.0])


def build_report() -> None:
    summary = read_csv(OUT_DIR / "finetune_summary_results.csv")
    seed_rows = read_csv(OUT_DIR / "finetune_seed_results.csv")
    config = json.loads((OUT_DIR / "finetune_config.json").read_text(encoding="utf-8"))
    device = seed_rows[0]["device"] if seed_rows else "CUDA GPU"

    doc = Document()
    configure_doc(doc)
    add_cover(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("基于 ImageNet 预训练模型的参数高效微调方法探索")
    set_east_asian_font(r, 18, True)

    add_heading(doc, "摘要", 1)
    add_para(
        doc,
        "大规模预训练模型已经成为视觉识别和自然语言处理中的重要基础模型，但对每个下游任务进行全量微调会带来较高的存储、训练和部署成本。"
        "参数高效微调通过冻结大部分预训练参数，只训练少量新增参数或特定参数子集，在保持迁移能力的同时显著降低训练开销。"
        "本文参考 AdaptFormer、SSF、LoRA 和 BitFit 四类代表性方法，基于 ImageNet 预训练 ResNet18 和 CIFAR-10 构建 CUDA 小样本实验，比较线性探测、BitFit、SSF、LoRA、Adapter 和全量微调六种策略。"
        "实验在 NVIDIA GeForce RTX 4060 Laptop GPU 上运行，每类使用 100 张训练图像和 100 张测试图像，重复两个随机种子。"
        "结果表明，BitFit 取得最高平均测试准确率 77.25%，但仅训练 0.0888% 的参数；全量微调训练 100% 参数，平均准确率为 72.80%。"
        "这说明在小样本迁移场景下，偏置项和任务头已经能够提供有效适配，全量更新未必带来更优泛化。"
    )
    add_para(doc, "关键词：模型微调；参数高效微调；BitFit；LoRA；SSF；Adapter；CUDA", first_line=False)

    add_heading(doc, "1 引言", 1)
    add_para(
        doc,
        "随着预训练模型规模不断扩大，直接为每个任务保存一份完整微调后的模型会造成明显的参数冗余。"
        "在视觉识别任务中，预训练骨干网络通常已经学习到较通用的边缘、纹理和语义表征，下游任务往往只需要对特征分布或决策边界进行较小调整。"
        "因此，如何在冻结主体参数的条件下完成有效适配，成为模型微调技术探索中的关键问题。"
    )
    add_para(
        doc,
        "已有研究从不同角度提出参数高效微调方案。AdaptFormer 在 Transformer 层中插入轻量适配器，以较少参数适配多种视觉识别任务；SSF 通过特征缩放和平移调节预训练特征；LoRA 将权重增量建模为低秩矩阵；BitFit 则进一步简化为只更新偏置项。"
        "这些方法共同体现出一个思想：下游任务的有效更新可能位于低维或少量参数子空间中。"
    )
    add_para(
        doc,
        "本文的主要贡献如下：第一，结合四篇参考文献整理了 Adapter、SSF、LoRA 和 BitFit 的核心机制；第二，使用 CUDA 实现了六种微调策略的统一对比实验；第三，根据真实实验结果分析不同参数高效策略在小样本图像分类场景下的效果、参数量和局限。"
    )

    add_heading(doc, "2 研究方法", 1)
    add_heading(doc, "2.1 参数高效微调思想", 2)
    add_para(
        doc,
        "全量微调会更新预训练模型的全部权重，表达能力最强，但训练参数量和过拟合风险也最高。"
        "参数高效微调通常冻结预训练骨干，只开放少量任务相关参数，使模型在保留原有表示能力的同时完成下游适配。"
        "本文将这些策略抽象为三个层次：只训练输出层的线性探测、训练少量原有参数的 BitFit、以及添加或重参数化增量模块的 SSF、LoRA 和 Adapter。"
    )
    add_heading(doc, "2.2 参考方法", 2)
    add_para(
        doc,
        "AdaptFormer 的启发来自 Adapter 思路：在主干网络中加入瓶颈结构，使输入特征先降维再升维，通过残差方式影响原模型输出。"
        "本文使用图像分类骨干 ResNet18，因此将该思想简化为特征向量后的瓶颈适配器。"
    )
    add_para(
        doc,
        "SSF 方法认为下游任务适配可以通过对预训练特征进行缩放和偏移实现。本文在 ResNet18 的全局池化特征上学习一组 scale 和 shift 参数，再连接分类头。"
    )
    add_para(
        doc,
        "LoRA 假设微调时的权重增量具有低秩结构，冻结原始权重并训练两个小矩阵的乘积。由于本实验的分类头是新初始化层，本文将 LoRA 简化应用在分类头增量上，这也构成后续结果分析中的一个重要局限。"
    )
    add_para(
        doc,
        "BitFit 只训练偏置项和任务相关分类层。该方法结构简单、参数量极低，适合作为参数高效微调的强基线。"
    )
    add_method_table(doc)
    add_caption(doc, "表 1 六种微调策略的实验实现")

    add_heading(doc, "3 实验与结果分析", 1)
    add_heading(doc, "3.1 实验设置", 2)
    add_para(
        doc,
        f"实验数据集为 CIFAR-10。为模拟小样本迁移场景，每个类别采样 {config['train_per_class']} 张训练图像和 {config['test_per_class']} 张测试图像。"
        f"模型采用 ImageNet 预训练 ResNet18，输入图像缩放到 {config['image_size']} x {config['image_size']}。"
        f"所有方法训练 {config['epochs']} 个 epoch，batch size 为 {config['batch_size']}，随机种子为 {config['seeds'][0]} 和 {config['seeds'][1]}。"
        f"实验设备为 {device}，训练过程中启用 CUDA 自动混合精度。"
    )
    add_para(
        doc,
        "评价指标包括测试准确率、测试损失、平均训练时间、总参数量、可训练参数量和可训练参数比例。"
        "其中 FULL 用作全量微调参考，LINEAR 用作最低成本迁移基线，其余方法对应参考文献中的参数高效微调思想。"
    )
    add_heading(doc, "3.2 实验结果", 2)
    add_result_table(doc, summary)
    add_caption(doc, "表 2 CUDA 小样本微调实验汇总结果")

    if CURVE_PNG.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(str(CURVE_PNG), width=Inches(5.9))
        add_caption(doc, "图 1 六种微调策略的测试准确率曲线")

    add_heading(doc, "3.3 结果分析", 2)
    add_para(
        doc,
        "从表 2 可以看出，BitFit 在本实验中表现最好，平均测试准确率达到 77.25%，比全量微调高 4.45 个百分点，同时可训练参数比例只有 0.0888%。"
        "这说明在 CIFAR-10 小样本迁移设置中，预训练 ResNet18 的主体特征已经具有较强泛化能力，适当调整偏置项和分类头即可显著改变任务决策边界。"
    )
    add_para(
        doc,
        "全量微调的平均准确率为 72.80%，低于 BitFit。其原因可能是训练样本较少、训练轮数较短，全量更新使大量预训练参数同时改变，容易受到小样本噪声影响。"
        "这与参数高效微调的动机一致：当下游数据不足时，减少可训练参数不仅节省计算，也可能改善泛化。"
    )
    add_para(
        doc,
        "LINEAR、SSF 和 ADAPTER 的平均准确率都在 71% 左右。其中 ADAPTER 的标准差最低，仅 0.57 个百分点，说明瓶颈适配器在两个随机种子下较稳定。"
        "SSF 没有超过线性探测，可能是因为本文只在最终特征上加入缩放和平移，而原论文通常在更细粒度层级中调节特征分布。"
    )
    add_para(
        doc,
        "LoRA 的平均准确率只有 14.25%，明显低于其他方法。该结果不应简单理解为 LoRA 方法无效，而应结合本实验实现分析：本文的 LoRA 作用于新初始化分类头，缺少原论文中对预训练权重矩阵进行低秩增量适配的条件；同时低秩约束使新分类头在短训练中表达能力不足。"
        "因此，这一结果反而提示 LoRA 更适合应用在已有预训练权重的线性投影层中，而不是只替代新任务头。"
    )

    add_heading(doc, "4 结论", 1)
    add_para(
        doc,
        "本文按照课程设计第四选题“模型微调技术探索”的要求，参考 AdaptFormer、SSF、LoRA 和 BitFit 四篇文献，完成了参数高效微调方法的整理、CUDA 实验实现和结果分析。"
        "实验表明，在 ImageNet 预训练 ResNet18 到 CIFAR-10 小样本分类的迁移任务中，BitFit 以极低可训练参数量获得最高准确率，说明简单的偏置更新可以成为强有力的参数高效微调基线。"
    )
    add_para(
        doc,
        "本实验也存在局限。第一，出于运行时间控制，只使用每类 100 张训练图像和 4 个 epoch，不能完全代表充分训练后的性能；第二，LoRA 和 AdaptFormer 都根据 ResNet18 结构做了简化实现，与原论文面向 Transformer 层的完整设置不同；第三，SSF 只作用于最终特征层，尚未覆盖更多中间层。"
        "后续可以扩展到 Vision Transformer、更多数据集和更多随机种子，并在预训练层内部实现 LoRA、SSF 和 Adapter，以更严格地复现参考文献设定。"
    )

    add_heading(doc, "5 参考文献", 1)
    references = [
        "Chen S., Ge C., Tong Z., Wang J., Song Y., Wang J., Luo P. AdaptFormer: Adapting Vision Transformers for Scalable Visual Recognition. NeurIPS, 2022.",
        "Lian D., Zhou D., Feng J., Wang X. Scaling & Shifting Your Features: A New Baseline for Efficient Model Tuning. NeurIPS, 2022.",
        "Hu E. J., Shen Y., Wallis P., Allen-Zhu Z., Li Y., Wang S., Wang L., Chen W. LoRA: Low-Rank Adaptation of Large Language Models. ICLR, 2022.",
        "Ben Zaken E., Goldberg Y., Ravfogel S. BitFit: Simple Parameter-efficient Fine-tuning for Transformer-based Masked Language-models. ACL, 2022.",
        "He K., Zhang X., Ren S., Sun J. Deep Residual Learning for Image Recognition. CVPR, 2016.",
        "Krizhevsky A. Learning Multiple Layers of Features from Tiny Images. 2009.",
    ]
    for idx, ref in enumerate(references, 1):
        add_para(doc, f"[{idx}] {ref}", first_line=False)

    add_heading(doc, "附录：复现实验文件", 1)
    add_para(doc, r"实验脚本：F:\Exploration\src\run_finetuning_cuda_experiment.py", first_line=False)
    add_para(doc, r"结果目录：F:\Exploration\outputs\finetune", first_line=False)
    add_para(doc, r"完成步骤：F:\Exploration\完成步骤.md", first_line=False)
    add_para(doc, r"运行方式：使用 CUDA 环境运行 python F:\Exploration\src\run_finetuning_cuda_experiment.py", first_line=False)

    doc.save(FINAL_DOCX)


if __name__ == "__main__":
    build_report()
